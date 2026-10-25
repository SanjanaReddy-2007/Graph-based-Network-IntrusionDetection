"""Assemble results/report.md and results/report.docx from the logged CSVs/figures.

Everything numeric is read from results/*.csv produced by real runs; if a file is missing the section says so
instead of inventing numbers. The H1/H2 verdict is a rule-based summary of the logged statistics: read it together
with the figures and edit the Discussion before submission."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
import pandas as pd

from src.utils import resolve_path

from . import report_text as T
from .common import log

NA = "not available (run the experiment first)"


def _read(out: Path, name: str) -> Optional[pd.DataFrame]:
    f = out / name
    return pd.read_csv(f) if f.exists() else None


def _pm(df: pd.DataFrame, col: str) -> pd.Series:
    return df[f"{col}_mean"].map("{:.3f}".format) + " ± " + df[f"{col}_std"].fillna(0).map("{:.3f}".format)


def tables(out: Path) -> Dict[str, Optional[pd.DataFrame]]:
    t: Dict[str, Optional[pd.DataFrame]] = {}
    s = _read(out, "protocols_summary.csv")
    if s is not None:
        t["protocols"] = pd.DataFrame({"Protocol": s.protocol, "Model": s.model, "Macro-F1": _pm(s, "macro_f1"),
                                       "PR-AUC": _pm(s, "pr_auc"), "Attack recall": _pm(s, "attack_recall"),
                                       "FPR": _pm(s, "fpr"), "MCC": _pm(s, "mcc"), "Accuracy*": _pm(s, "accuracy"),
                                       "Weighted-F1*": _pm(s, "weighted_f1"), "Seeds": s.n_seeds})
    g = _read(out, "protocols_gaps.csv")
    if g is not None:
        t["protocol_gaps"] = pd.DataFrame({
            "Protocol": g.protocol, "A - B": g.A + " - " + g.B,
            "Δ macro-F1": g.macro_f1_mean_gap.map("{:+.3f}".format),
            "95% CI (boot.)": g.macro_f1_ci_lo.map("[{:+.3f}".format) + ", " + g.macro_f1_ci_hi.map("{:+.3f}]".format),
            "Δ PR-AUC": g.pr_auc_mean_gap.map("{:+.3f}".format),
            "t-test p": g.macro_f1_t_pvalue.map("{:.3g}".format), "Wilcoxon p": g.macro_f1_wilcoxon_pvalue.map("{:.3g}".format)})
    o = _read(out, "protocols_overlap.csv")
    if o is not None:
        cols = {"ref_host_containment": "Host contain.", "ref_host_jaccard": "Host Jaccard", "ref_one_hop_mean": "1-hop ovl (ref)",
                "ref_two_hop_train_frac": "2-hop ovl (ref)", "ref_test_cross_edge_ratio": "Cross-edge (test)",
                "ref_seen_host_test_nbr": "Seen-host test nbr", "ref_edge_homophily": "Edge homophily",
                "ref_flow5tuple_seen": "5-tuple seen", "vis_one_hop_mean": "1-hop ovl (visible)"}
        m = o.groupby("protocol")[[c for c in cols if c in o]].mean().rename(columns=cols).round(3).reset_index()
        t["overlap"] = m.rename(columns={"protocol": "Protocol"})
    sg = _read(out, "sweep_gaps.csv")
    if sg is not None:
        d = pd.DataFrame({"Knob": sg.knob, "Target level": sg.level,
                          "Achieved host contain.": sg.get("achieved_ref_host_containment", np.nan).round(3),
                          "Achieved 1-hop (ref)": sg.get("achieved_ref_one_hop_mean", np.nan).round(3),
                          "Achieved 1-hop (visible)": sg.get("achieved_vis_one_hop_mean", np.nan).round(3),
                          "Δ macro-F1 (SAGE-MLP)": sg.macro_f1_mean_gap.map("{:+.3f}".format),
                          "95% CI": sg.macro_f1_ci_lo.map("[{:+.3f}".format) + ", " + sg.macro_f1_ci_hi.map("{:+.3f}]".format),
                          "Δ PR-AUC": sg.pr_auc_mean_gap.map("{:+.3f}".format)})
        t["sweep"] = d
    sc = _read(out, "sweep_correlations.csv")
    if sc is not None:
        c = sc[sc.gap_metric == "gap_macro_f1"]
        t["correlations"] = pd.DataFrame({"Scope": c.scope, "Overlap measure": c.overlap_measure, "n": c.n.astype(int),
                                          "Spearman": c.spearman.map("{:+.2f}".format), "p": c.spearman_p.map("{:.3g}".format),
                                          "Pearson": c.pearson.map("{:+.2f}".format), "p ": c.pearson_p.map("{:.3g}".format)})
    ag = _read(out, "ablation_gaps.csv")
    if ag is not None:
        t["ablation"] = pd.DataFrame({"Protocol": ag.protocol, "Comparison": ag.comparison,
                                      "Δ macro-F1": ag.macro_f1_mean_gap.map("{:+.3f}".format),
                                      "95% CI": ag.macro_f1_ci_lo.map("[{:+.3f}".format) + ", " + ag.macro_f1_ci_hi.map("{:+.3f}]".format),
                                      "t-test p": ag.macro_f1_t_pvalue.map("{:.3g}".format)})
    asum = _read(out, "ablation_summary.csv")
    if asum is not None:
        t["ablation_summary"] = pd.DataFrame({"Protocol": asum.protocol, "Model / topology": asum.model_topology,
                                              "Macro-F1": _pm(asum, "macro_f1"), "PR-AUC": _pm(asum, "pr_auc"),
                                              "Attack recall": _pm(asum, "attack_recall")})
    pc = out / "param_counts.json"
    if pc.exists():
        p = json.loads(pc.read_text())
        t["params"] = pd.DataFrame([{"Input features": p["n_features"], "GraphSAGE hidden": p["sage_hidden"],
                                     "GraphSAGE params": p["sage_params"], "MLP hidden": p["mlp_hidden"],
                                     "MLP params": p["mlp_params"], "Relative diff.": f"{p['rel_diff']:.2%}"}])
    return t


def verdict(out: Path) -> List[str]:
    """Rule-based reading of the logged statistics (see docstring). Returns sentences."""
    pg, sg, sc = _read(out, "protocols_gaps.csv"), _read(out, "sweep_gaps.csv"), _read(out, "sweep_correlations.csv")
    if pg is None or sg is None:
        return ["The verdict cannot be generated because the protocol and/or overlap-sweep results are missing."]
    L: List[str] = []
    sm = pg[(pg.A == "sage") & (pg.B == "mlp")].set_index("protocol")
    for p in sm.index:
        r = sm.loc[p]
        L.append(f"Under {p} the GraphSAGE-MLP macro-F1 gap is {r.macro_f1_mean_gap:+.3f} "
                 f"(95% CI [{r.macro_f1_ci_lo:+.3f}, {r.macro_f1_ci_hi:+.3f}]).")
    keep = {}
    if "P1" in sm.index and "P2" in sm.index:
        d = sm.loc["P1", "macro_f1_mean_gap"] - sm.loc["P2", "macro_f1_mean_gap"]
        L.append(f"Removing cross-split structure (P1 → P2, same node split) changes the gap by {-d:+.3f}.")
    low = sg.sort_values("level").groupby("knob").first()
    for knob, r in low.iterrows():
        pos = r.macro_f1_ci_lo > 0
        L.append(f"At the lowest {knob}-overlap level (target {r.level:g}) the gap is {r.macro_f1_mean_gap:+.3f} "
                 f"[{r.macro_f1_ci_lo:+.3f}, {r.macro_f1_ci_hi:+.3f}] - {'its CI excludes zero' if pos else 'its CI includes zero (or is negative)'}.")
        keep[knob] = pos
    sig_pos = False
    if sc is not None:
        c = sc[(sc.gap_metric == "gap_macro_f1") & (sc.scope.isin(["host", "edge"])) & (sc.overlap_measure != "level")]
        for _, r in c.iterrows():
            if not np.isnan(r.spearman):
                L.append(f"Correlation of the gap with {r.overlap_measure} ({r.scope} knob, n={int(r.n)} runs): "
                         f"Spearman {r.spearman:+.2f} (p={r.spearman_p:.3g}).")
        sig_pos = bool(((c.spearman > 0) & (c.spearman_p < 0.05)).any())
    inductive_pos = bool(sm.loc[[p for p in ("P2", "P3") if p in sm.index], "macro_f1_ci_lo"].gt(0).all()) if len(sm) else False
    n_keep, n_knobs = sum(keep.values()), len(keep)
    if inductive_pos and n_keep == n_knobs and not sig_pos:
        v = ("Taken together, these statistics are consistent with H2: GraphSAGE retains a positive advantage under leakage-safe and "
             "low-overlap evaluation (CIs exclude zero) and the gap does not shrink significantly with overlap.")
    elif sig_pos and n_keep == 0 and not inductive_pos:
        v = ("Taken together, these statistics are consistent with H1: the gap grows with overlap and is not clearly "
             "distinguishable from zero at the lowest overlap levels or under the leakage-safe protocols.")
    elif sig_pos or inductive_pos or n_keep > 0:
        v = ("The evidence is MIXED: part of the advantage depends on overlap (H1) while part persists (H2) - for example, the gap may shrink "
             "under one overlap knob or protocol but remain positive under another. Neither hypothesis is supported in its pure form; "
             "see the per-protocol and per-knob numbers above.")
    else:
        v = "The statistics do not clearly favour either hypothesis (inconclusive): the gap is not distinguishable from zero anywhere."
    L.append(v)
    return L


def build_markdown(cfg: dict, out: Path, tb: Dict[str, Optional[pd.DataFrame]], vd: List[str]) -> str:
    def md(df):
        if df is None:
            return f"*{NA}*\n"
        return df.to_markdown(index=False) + "\n" if hasattr(df, "to_markdown") and _has_tabulate() else _simple_md(df)

    s = [f"# {T.TITLE}", f"*{T.AUTHORS}*", "", "## Abstract", T.ABSTRACT, "", "## 1. Introduction and Motivation", *T.INTRO,
         "", "## 2. Related Work", T.RELATED_INTRO, "", "**Table 1.** Comparison with related work.", "",
         _simple_md(pd.DataFrame(T.RELATED_ROWS, columns=T.RELATED_HEAD)), "", "## 3. Network-Security Perspective"]
    for h, p in T.NETWORK_ANGLE:
        s += [f"**{h}.** {p}", ""]
    s += ["## 4. Methodology"]
    for h, p in T.METHOD:
        s += [f"**{h}.** {p}", ""]
    s += ["## 5. Experimental Setup",
          f"Seeds: {cfg['seeds']}. Graph: window W = {cfg['graph']['window_seconds']} s, {cfg['graph']['neighbors_per_type']} predecessors per entity type, "
          f"degree cap {cfg['graph']['max_degree']}, entity types {cfg['graph']['entity_types']}. GraphSAGE hidden width {cfg['model']['hidden']}, dropout {cfg['model']['dropout']}, "
          f"Adam lr {cfg['train']['lr']}, weight decay {cfg['train']['weight_decay']}, max {cfg['train']['epochs']} epochs, patience {cfg['train']['patience']}. "
          f"Random Forest: {cfg['rf']['n_estimators']} trees, balanced class weights. Temporal protocol: train on {cfg['protocols']['temporal']['train_days']}, "
          f"test on {cfg['protocols']['temporal']['test_days']}. Sub-sampling: {cfg['subsample']}.", "",
          "**Table 2.** Capacity matching.", "", md(tb.get("params")), ""]
    eda = out / "eda_summary.json"
    if eda.exists():
        e = json.loads(eda.read_text())
        s += [f"Processed data: {e['n_flows']} flows, {e['n_features']} features, {e['n_ips']} IPs, attack fraction {e['attack_fraction']:.3f}; "
              f"flow graph mean degree {e['mean_degree']:.2f}, isolated nodes {e['isolated_fraction']:.1%}, edge label homophily {e['edge_homophily']:.3f}.", ""]
    s += ["## 6. Results", "### 6.1 Performance under the three protocols", "**Table 3.** Mean ± std over seeds (*accuracy and weighted-F1 are shown only for comparison with prior work).", "",
          md(tb.get("protocols")), "", "![](fig_protocols.png)", "", "**Table 4.** Paired gaps between models.", "", md(tb.get("protocol_gaps")), "",
          "### 6.2 Measured structural overlap", "**Table 5.** Overlap metrics per protocol (mean over seeds; 'ref' = reference graph, 'visible' = graph used by the model).", "",
          md(tb.get("overlap")), "", "![](fig_nbr_overlap_dist.png)", "",
          "### 6.3 Controlled overlap variation", "**Table 6.** Achieved overlap and GraphSAGE-MLP gap.", "", md(tb.get("sweep")), "",
          "![](fig_overlap_gap.png)", "", "**Table 7.** Correlation between overlap measures and the macro-F1 gap (per run).", "", md(tb.get("correlations")), "",
          "### 6.4 Graph-structure ablation", "**Table 8.** Macro-F1 by topology.", "", md(tb.get("ablation_summary")), "", "![](fig_ablation.png)", "",
          "**Table 9.** Paired differences (original topology minus disrupted; sanity check: GraphSAGE without edges minus MLP).", "", md(tb.get("ablation")), "",
          "### 6.5 Per-attack-type recall", "![](fig_recall_heatmap.png)", "",
          "## 7. Discussion", "*Auto-generated summary of the logged statistics - review and edit before submission:*", ""]
    s += [f"- {x}" for x in vd]
    s += ["", "## 8. Threats to Validity"] + [f"- {x}" for x in T.THREATS]
    s += ["", "## 9. Conclusion",
          "We provided an overlap-aware, controlled evaluation framework for graph-based NIDS: explicit train/test overlap metrics, three protocols with leakage assertions, "
          "controlled overlap variation, capacity-matched baselines and a topology ablation. The conclusion about H1 versus H2 must be taken from Section 7 and the tables above, "
          "which are generated from the executed runs.", "", "## References"] + [f"{r}  " for r in T.REFERENCES]
    return "\n".join(s)


def _has_tabulate() -> bool:
    try:
        import tabulate  # noqa: F401
        return True
    except ImportError:
        return False


def _simple_md(df: pd.DataFrame) -> str:
    cols = [str(c) for c in df.columns]
    lines = ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(str(v) for v in r.values) + " |")
    return "\n".join(lines) + "\n"


def build_docx(cfg: dict, out: Path, tb, vd: List[str], path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Inches, Pt
    except ImportError:
        log.warning("python-docx not installed: skipping report.docx (pip install python-docx)")
        return
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.9)

    def table(df, small=True):
        if df is None:
            doc.add_paragraph(NA).italic = True
            return
        t = doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(df.columns):
            t.rows[0].cells[i].text = str(c)
        for _, r in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(r.values):
                cells[i].text = str(v)
        if small:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(7.5)
        doc.add_paragraph()

    def pic(name, w=6.2):
        if (out / name).exists():
            doc.add_picture(str(out / name), width=Inches(w))

    doc.add_heading(T.TITLE, 0)
    doc.add_paragraph(T.AUTHORS).italic = True
    doc.add_heading("Abstract", 1); doc.add_paragraph(T.ABSTRACT)
    doc.add_heading("1. Introduction and Motivation", 1)
    for p in T.INTRO: doc.add_paragraph(p)
    doc.add_heading("2. Related Work", 1); doc.add_paragraph(T.RELATED_INTRO)
    doc.add_paragraph("Table 1. Comparison with related work.")
    table(pd.DataFrame(T.RELATED_ROWS, columns=T.RELATED_HEAD))
    doc.add_heading("3. Network-Security Perspective", 1)
    for h, p in T.NETWORK_ANGLE:
        para = doc.add_paragraph(); para.add_run(h + ". ").bold = True; para.add_run(p)
    doc.add_heading("4. Methodology", 1)
    for h, p in T.METHOD:
        para = doc.add_paragraph(); para.add_run(h + ". ").bold = True; para.add_run(p)
    doc.add_heading("5. Experimental Setup", 1)
    doc.add_paragraph(f"Seeds: {cfg['seeds']}. Graph: W = {cfg['graph']['window_seconds']} s, {cfg['graph']['neighbors_per_type']} predecessors per entity type, "
                      f"degree cap {cfg['graph']['max_degree']}. GraphSAGE hidden {cfg['model']['hidden']}, dropout {cfg['model']['dropout']}, Adam lr {cfg['train']['lr']}, "
                      f"max {cfg['train']['epochs']} epochs, patience {cfg['train']['patience']}. RF: {cfg['rf']['n_estimators']} trees. Sub-sampling: {cfg['subsample']}.")
    doc.add_paragraph("Table 2. Capacity matching."); table(tb.get("params"), False)
    doc.add_heading("6. Results", 1)
    doc.add_heading("6.1 Performance under the three protocols", 2)
    doc.add_paragraph("Table 3. Mean ± std over seeds (*for comparison with prior work only).")
    table(tb.get("protocols")); pic("fig_protocols.png")
    doc.add_paragraph("Table 4. Paired gaps between models."); table(tb.get("protocol_gaps"))
    doc.add_heading("6.2 Measured structural overlap", 2)
    doc.add_paragraph("Table 5. Overlap metrics per protocol."); table(tb.get("overlap")); pic("fig_nbr_overlap_dist.png")
    doc.add_heading("6.3 Controlled overlap variation", 2)
    doc.add_paragraph("Table 6. Achieved overlap and GraphSAGE-MLP gap."); table(tb.get("sweep")); pic("fig_overlap_gap.png")
    doc.add_paragraph("Table 7. Correlation between overlap measures and the macro-F1 gap."); table(tb.get("correlations"))
    doc.add_heading("6.4 Graph-structure ablation", 2)
    doc.add_paragraph("Table 8. Macro-F1 by topology."); table(tb.get("ablation_summary")); pic("fig_ablation.png")
    doc.add_paragraph("Table 9. Paired differences."); table(tb.get("ablation"))
    doc.add_heading("6.5 Per-attack-type recall", 2); pic("fig_recall_heatmap.png", 5.5)
    doc.add_heading("7. Discussion", 1)
    doc.add_paragraph("Auto-generated summary of the logged statistics - review and edit before submission:").italic = True
    for x in vd: doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("8. Threats to Validity", 1)
    for x in T.THREATS: doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("9. Conclusion", 1)
    doc.add_paragraph("We provided an overlap-aware, controlled evaluation framework for graph-based NIDS. The conclusion about H1 versus H2 must be taken from Section 7 and the tables, which are generated from the executed runs.")
    doc.add_heading("References", 1)
    for r in T.REFERENCES: doc.add_paragraph(r)
    doc.save(path)


def main(cfg: dict) -> None:
    out = resolve_path(cfg, "results_dir")
    tb = tables(out)
    vd = verdict(out)
    (out / "report.md").write_text(build_markdown(cfg, out, tb, vd), encoding="utf-8")
    build_docx(cfg, out, tb, vd, out / "report.docx")
    for k, v in tb.items():
        if v is not None:
            v.to_csv(out / f"table_{k}.csv", index=False)
    log.info("Report written: %s and %s", out / "report.md", out / "report.docx")
